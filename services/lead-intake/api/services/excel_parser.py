"""
api/services/excel_parser.py — Excel parsing for bulk lead import.

Uses openpyxl in read-only/streaming mode rather than pandas: bulk-import
files are read once, header-mapped, and discarded — there's no need for a
full DataFrame, and read_only mode keeps memory flat for large spreadsheets.
"""
from __future__ import annotations

import io
from typing import Any, Iterator, Optional

import openpyxl
from fastapi import HTTPException, status


def _load_sheet(file_bytes: bytes, sheet_name: Optional[str]):
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={
                "error": "INVALID_FILE",
                "message": f"Could not read the uploaded file as an Excel workbook: {exc}",
                "details": [],
            },
        )

    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail={
                    "error": "INVALID_FILE",
                    "message": f"Sheet '{sheet_name}' not found. Available: {workbook.sheetnames}",
                    "details": [],
                },
            )
        return workbook[sheet_name]
    return workbook.worksheets[0]


def parse_header_row(
    file_bytes: bytes,
    sheet_name: Optional[str] = None,
    header_row_index: int = 0,
) -> list[str]:
    """
    Read just the header row (0-based index) and return the non-empty cell
    values as strings. Used to verify a client-submitted column mapping still
    matches the file's real headers before trusting it.
    """
    sheet = _load_sheet(file_bytes, sheet_name)
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == header_row_index:
            return [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
        if i > header_row_index:
            break

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail={
            "error": "INVALID_FILE",
            "message": f"Sheet has no row at headerRowIndex={header_row_index}.",
            "details": [],
        },
    )


def iter_rows(
    file_bytes: bytes,
    sheet_name: Optional[str] = None,
    header_row_index: int = 0,
) -> Iterator[dict[str, Any]]:
    """
    Stream data rows (after the header row) as {header: value} dicts.
    Skips fully-blank rows. Uses data_only=True so formula cells yield their
    last computed value, not the formula text.
    """
    sheet = _load_sheet(file_bytes, sheet_name)
    headers: Optional[list[str]] = None

    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i < header_row_index:
            continue
        if i == header_row_index:
            headers = [str(cell).strip() if cell is not None else "" for cell in row]
            continue

        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue  # skip blank rows

        row_dict = {
            headers[col_idx]: value
            for col_idx, value in enumerate(row)
            if col_idx < len(headers) and headers[col_idx]
        }
        yield row_dict
