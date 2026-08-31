"""
tests/unit/test_excel_parser.py — Tests for services/excel_parser.py.

Covers header-row parsing (with title-row offsets), sheet-name selection,
row streaming with blank rows, and rejection of non-Excel / corrupt files.
"""
from __future__ import annotations

import io
import os
import sys
from datetime import date, datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "../../api"))
os.environ.setdefault("FIREBASE_PROJECT_ID", "test-project")

import openpyxl
import pytest
from fastapi import HTTPException

from services.excel_parser import iter_rows, parse_header_row


def _build_xlsx(rows: list[list], sheet_name: str = "Sheet1") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParseHeaderRow:
    def test_reads_simple_header_row(self):
        data = _build_xlsx([
            ["First Name", "Last Name", "Email"],
            ["John", "Doe", "john@example.com"],
        ])
        headers = parse_header_row(data)
        assert headers == ["First Name", "Last Name", "Email"]

    def test_reads_header_at_offset(self):
        """Some spreadsheets have a title row above the real headers."""
        data = _build_xlsx([
            ["Q1 2026 Lead Import"],
            ["First Name", "Last Name", "Email"],
            ["John", "Doe", "john@example.com"],
        ])
        headers = parse_header_row(data, header_row_index=1)
        assert headers == ["First Name", "Last Name", "Email"]

    def test_ignores_blank_trailing_cells(self):
        data = _build_xlsx([["First Name", "Last Name", None, None]])
        headers = parse_header_row(data)
        assert headers == ["First Name", "Last Name"]

    def test_reads_named_sheet(self):
        data = _build_xlsx([["A", "B"]], sheet_name="Leads")
        headers = parse_header_row(data, sheet_name="Leads")
        assert headers == ["A", "B"]

    def test_missing_sheet_raises_400(self):
        data = _build_xlsx([["A", "B"]], sheet_name="Leads")
        with pytest.raises(HTTPException) as exc_info:
            parse_header_row(data, sheet_name="DoesNotExist")
        assert exc_info.value.status_code == 400

    def test_header_row_index_beyond_data_raises_400(self):
        data = _build_xlsx([["A", "B"]])
        with pytest.raises(HTTPException) as exc_info:
            parse_header_row(data, header_row_index=5)
        assert exc_info.value.status_code == 400

    def test_corrupt_file_raises_400(self):
        with pytest.raises(HTTPException) as exc_info:
            parse_header_row(b"not a real xlsx file")
        assert exc_info.value.status_code == 400


class TestIterRows:
    def test_yields_row_dicts_keyed_by_header(self):
        data = _build_xlsx([
            ["First Name", "Last Name"],
            ["John", "Doe"],
            ["Jane", "Smith"],
        ])
        rows = list(iter_rows(data))
        assert rows == [
            {"First Name": "John", "Last Name": "Doe"},
            {"First Name": "Jane", "Last Name": "Smith"},
        ]

    def test_skips_fully_blank_rows(self):
        data = _build_xlsx([
            ["First Name", "Last Name"],
            ["John", "Doe"],
            [None, None],
            ["Jane", "Smith"],
        ])
        rows = list(iter_rows(data))
        assert len(rows) == 2
        assert rows[0]["First Name"] == "John"
        assert rows[1]["First Name"] == "Jane"

    def test_respects_header_row_index_offset(self):
        data = _build_xlsx([
            ["Title Row"],
            ["First Name", "Last Name"],
            ["John", "Doe"],
        ])
        rows = list(iter_rows(data, header_row_index=1))
        assert rows == [{"First Name": "John", "Last Name": "Doe"}]

    def test_preserves_date_cell_types(self):
        data = _build_xlsx([
            ["First Name", "DOB"],
            ["John", date(2001, 9, 11)],
        ])
        rows = list(iter_rows(data))
        dob = rows[0]["DOB"]
        assert isinstance(dob, (date, datetime))
        assert dob.year == 2001 and dob.month == 9 and dob.day == 11

    def test_no_data_rows_yields_empty_iterator(self):
        data = _build_xlsx([["First Name", "Last Name"]])
        assert list(iter_rows(data)) == []
